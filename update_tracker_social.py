"""Refresh Viral Formats tab with competitor-social research + update log."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

P = "/home/hatch/workspace/junkmove-site/Junk-Move-Pros-Competitor-Tracker.xlsx"
wb = openpyxl.load_workbook(P)

HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
HDR_FILL = PatternFill("solid", fgColor="1F4E5F")
SUB_FILL = PatternFill("solid", fgColor="D9E8EC")
THIN = Border(*(Side(style="thin", color="B0B0B0") for _ in ("left", "right", "top", "bottom")))
WRAP = Alignment(wrap_text=True, vertical="top")

wb.remove(wb["Viral Formats"])
ws = wb.create_sheet("Viral Formats", 5)

headers = ["#", "Format", "Hook formula (first 3 seconds)", "Example title/caption",
           "Platform", "Faceless?", "Proof it works", "Why we beat franchises at it"]
rows = [
 [1, "60-Second Point-It-Out Challenge", "We're in [TOWN] - whatever you can point to in 60 seconds, we haul FREE",
  "60-Second Cleanout Challenge in Ellenville, NY - she pointed at HOW much?!", "TikTok/Reels", "Yes",
  "1-800-GOT-JUNK? format; IG 110 likes, uniformly positive comments", "Corporate version feels staged; real crew + real neighbors = 10x believability"],
 [2, "After-First Transformation", "Open on the clean AFTER for 2 sec, smash-cut to the BEFORE",
  "This Kingston garage 6 hours ago vs now", "TikTok/Reels/Shorts", "Yes",
  "Universal viral mechanic; junk is the perfect before/after subject", "Raw phone footage beats their polished versions"],
 [3, "The Honesty Pledge", "If we misquote your job, we eat the difference. Here's why.",
  "We quoted $X. It took longer. We charged $X anyway. That's the policy.", "TikTok/Reels", "Yes",
  "@318.junk.removal: 18K plays, 1,033 likes, 75 shares", "Franchise comment sections call them liars; a local with a whiteboard is believed"],
 [4, "Compassion Cleanout", "She couldn't clear her porch after surgery. So we showed up.",
  "Helping a Newburgh widow get her home back - no charge", "TikTok/Reels", "Yes",
  "Junk Wranglers: 3.4K likes, 584 comments (deepest engagement in dataset)", "Corporate 'community' posts feel like PR; local crew helping a neighbor is town news"],
 [5, "One-Line Curiosity Caption", "Mundane footage + one provocative line; let comments ask",
  "Landlord called. Tenant left THIS.", "TikTok/Reels", "Yes",
  "@powerhaulers: 7,087 likes on 'Eviction clean out'", "Costs nothing; comments do the distribution work"],
 [6, "Money Transparency", "What a full truckload actually costs us vs what we charge",
  "We made $___ on this cleanout - full breakdown", "TikTok/YouTube", "Yes",
  "Koerner Office 3.3K likes; Junk Teens 400K followers on money angle", "RANGES ONLY - never the internal floor or markup tiers"],
 [7, "Franchise Comparison, Local Edition", "[Franchise] quoted them $X over the phone. Here's what the job actually took.",
  "They called the 1-800 number first. Then they called us.", "YouTube", "Yes",
  "Tyler's Hauling (Rochester) owns it as a blog; nobody owns the HV video", "Factual only - no trademark abuse"],
 [8, "Day-in-the-Life POV", "Chest-mount, timestamps, real jobs, one real moment per video",
  "POV: 3 cleanouts, 1 dump run, 47 miles - a Tuesday in the Hudson Valley", "TikTok/YouTube", "Yes",
  "College Hunks Tampa ride-along 260 likes; HV locals do it at 3-11 likes (hooks missing)", "Format is proven; we bring better hooks"],
 [9, "Satisfying Load ASMR", "No talking, natural sound, fast cuts; text overlay carries story",
  "Oddly satisfying: 2 tons of garage junk in 40 seconds", "TikTok/Reels", "Yes",
  "The niche's background radiation - always performs", "Zero production cost; film on every job"],
 [10, "Where-It-Went Proof", "Your couch didn't go to the dump. Watch where it went.",
  "This estate cleanout furnished 3 families - here's the donation receipt", "TikTok/Reels", "Yes",
  "Franchises get accused of trashing donations; film the drop", "Donation receipts are free trust content"],
]

for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = HDR_FONT; cell.fill = HDR_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center"); cell.border = THIN
for r, row in enumerate(rows, 2):
    for c, v in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.alignment = WRAP; cell.border = THIN
        if r % 2 == 0:
            cell.fill = SUB_FILL
widths = [5, 30, 44, 44, 16, 10, 40, 44]
for c, w in enumerate(widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False

log = wb["Update Log"]
log.append(["2026-10-01", "Viral Formats tab refreshed with competitor-social research (10 formats + proof + why-we-win)",
            "Ziggy", "Source: research/competitor-social-2026-09-30.md; key gap: zero HV junk companies on TikTok"])
for cell in log[log.max_row]:
    cell.alignment = WRAP; cell.border = THIN

wb.save(P)
print("saved", P)
